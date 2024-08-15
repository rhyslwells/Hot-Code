import csv
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
import os

# Function to read parameters from the CSV file
def read_parameters_from_csv(csv_file):
    params = []
    with open(csv_file, mode='r') as file:
        reader = csv.DictReader(file)
        for row in reader:
            params.append(row)
    return params

# Function to map color names to RGBColor objects
def get_color(color_name):
    colors = {
        'red': RGBColor(255, 0, 0),
        'blue': RGBColor(0, 0, 255),
        'green': RGBColor(0, 128, 0),
        'orange': RGBColor(255, 165, 0),
        # Add more colors as needed
    }
    return colors.get(color_name.lower(), RGBColor(0, 0, 0))  # Default to black if color not found

# Function to replace placeholders and apply formatting in a .docx file
def create_docx_with_replacements(template_docx_path, params):
    for param in params:
        # Load the template document
        doc = Document(template_docx_path)
        
        # Replace placeholders and apply formatting
        for paragraph in doc.paragraphs:
            if 'xxx1' in paragraph.text:
                replace_text_with_formatting(paragraph, 'xxx1', param["xxx1_text"], int(param["xxx1_font_size"]), get_color(param["xxx1_color"]))
            if 'xxx2' in paragraph.text:
                replace_text_with_formatting(paragraph, 'xxx2', param["xxx2_text"], int(param["xxx2_font_size"]), get_color(param["xxx2_color"]))
            # Add more replacements as needed (xxx3, xxx4, etc.)

        # Save the new document
        output_filename = param["output_file"]
        doc.save(output_filename)
        print(f"{output_filename} created successfully.")

# Function to replace text and apply formatting
def replace_text_with_formatting(paragraph, placeholder, new_text, font_size, font_color):
    # Replace the placeholder while preserving formatting
    if placeholder in paragraph.text:
        inline = paragraph.runs
        for run in inline:
            if placeholder in run.text:
                run.text = run.text.replace(placeholder, new_text)
                run.font.size = Pt(font_size)
                run.font.color.rgb = font_color

# Read parameters from CSV
params = read_parameters_from_csv('parameters.csv')

# Use the function with the template DOCX in the "docs" folder and parameters
template_path = os.path.join('docs', 'template.docx')
create_docx_with_replacements(template_path, params)
